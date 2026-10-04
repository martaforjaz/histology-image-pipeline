"""Integrity checks must catch incomplete exports without modifying images."""
from pathlib import Path
import importlib.util
import struct
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

import numpy as np
import tifffile
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import validate_downsampling as validation


class OutputValidationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / '2x').mkdir()
        (self.root / '40x').mkdir()
        self.source = self.root / 'sample.ndpi'
        self.source.touch()
        self.path = self.root / '2x' / 'sample.tif'
        self.data = np.arange(48*64*3, dtype='uint8').reshape(48,64,3)

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, **kwargs):
        tifffile.imwrite(self.path, self.data, photometric='rgb', resolution=(2000,2000),
                         resolutionunit='CENTIMETER', **kwargs)

    def check(self):
        job = next(j for j in validation.inventory(self.root) if j['folder']=='2x')
        with patch.object(validation, 'source_geometry', return_value=((640,480),(.5,.5))):
            return validation.validate_file(job)

    def test_valid_stripped_and_tiled_exports_pass_without_modification(self):
        for options in ({}, {'tile':(16,16), 'compression':'deflate'}):
            self.write(**options)
            before = self.path.read_bytes()
            result = self.check()
            self.assertEqual(result['status'], 'PASS', result['reason'])
            self.assertTrue(result['full_decode'])
            self.assertEqual(self.path.read_bytes(), before)

    def test_truncated_pixel_data_fails(self):
        self.write()
        data = self.path.read_bytes()
        self.path.write_bytes(data[:-100])
        result = self.check()
        self.assertEqual(result['status'], 'FAIL')
        self.assertIn('beyond', result['reason'])

    def test_corrupt_compressed_tile_fails(self):
        self.write(tile=(16,16), compression='deflate')
        with tifffile.TiffFile(self.path) as tif:
            offset = tif.pages[0].dataoffsets[-1]
        with self.path.open('r+b') as f:
            f.seek(offset)
            f.write(b'BAD CODEC DATA')
        self.assertEqual(self.check()['status'], 'FAIL')

    def test_wrong_geometry_fails_even_when_all_pixels_decode(self):
        self.data = self.data[:, :32]
        self.write()
        result = self.check()
        self.assertEqual(result['status'], 'FAIL')
        self.assertTrue(result['full_decode'])
        self.assertIn('Wrong dimensions', result['reason'])

    def test_wrong_calibration_fails(self):
        tifffile.imwrite(self.path, self.data, resolution=(4000,4000), resolutionunit='CENTIMETER')
        result = self.check()
        self.assertEqual(result['status'], 'FAIL')
        self.assertIn('Wrong pixel spacing', result['reason'])

    def test_missing_and_partial_outputs_are_not_passed(self):
        self.write()
        self.path.rename(self.path.with_suffix('.tif.part'))
        jobs = validation.inventory(self.root)
        self.assertEqual(len(jobs), 2)
        results = [validation.validate_file(job) for job in jobs]
        self.assertTrue(all(r['status']=='FAIL' for r in results))
        self.assertTrue(any('Missing' in r['reason'] for r in results))
        self.assertTrue(any('Unfinished' in r['reason'] for r in results))

    def test_no_raw_geometry_is_review_not_green(self):
        self.write()
        job = next(j for j in validation.inventory(self.root) if j['folder']=='2x')
        with patch.object(validation, 'source_geometry', side_effect=ValueError('missing calibration')):
            result = validation.validate_file(job)
        self.assertEqual(result['status'], 'REVIEW')
        self.assertTrue(result['full_decode'])

    def test_blank_image_requires_review(self):
        self.data.fill(255)
        self.write()
        self.assertEqual(self.check()['status'], 'REVIEW')

    def test_network_read_error_is_unverified_not_corrupt(self):
        self.write()
        with patch.object(validation, 'check_tiff_directories', side_effect=OSError('Network read failed')):
            result = self.check()
        self.assertEqual(result['status'], 'REVIEW')
        self.assertFalse(result['full_decode'])

    def test_cyclic_directory_is_rejected(self):
        self.write()
        with self.path.open('r+b') as f:
            f.seek(4)
            offset = struct.unpack('<I', f.read(4))[0]
            f.seek(offset)
            count = struct.unpack('<H', f.read(2))[0]
            f.seek(offset + 2 + count*12)
            f.write(struct.pack('<I', offset))
        with self.assertRaisesRegex(ValueError, 'cyclic'):
            validation.check_tiff_directories(self.path)

    def test_timeout_is_review_not_corruption(self):
        self.write()
        job = next(j for j in validation.inventory(self.root) if j['folder']=='2x')
        worker = Mock()
        worker.communicate.side_effect = [validation.subprocess.TimeoutExpired('worker', 1), ('', '')]
        with patch.object(validation.subprocess, 'Popen', return_value=worker), patch.object(validation.subprocess, 'run'):
            result = validation.run_isolated(job, 1)
        self.assertEqual(result['status'], 'REVIEW')
        self.assertFalse(result['full_decode'])

    def test_valid_ome_pyramid_checks_every_page(self):
        import WSI2OMEtif_All_file_types as converter
        data = np.tile(self.data, (3,3,1))
        converter.save_ome_tif(Image.fromarray(data), str(self.root), '40x', 'sample', .25)
        job = next(j for j in validation.inventory(self.root) if j['folder']=='40x')
        with patch.object(validation, 'source_geometry', return_value=((192,144),(.25,.25))):
            result = validation.validate_file(job)
        self.assertEqual(result['status'], 'PASS', result['reason'])
        self.assertEqual(result['pages_checked'], 8)

    def test_duplicate_outputs_are_flagged(self):
        self.write()
        self.path.with_suffix('.ome.tif').write_bytes(self.path.read_bytes())
        self.assertEqual(self.check()['status'], 'FAIL')

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'Optional validation-report dependency')
    def test_portable_excel_export_status_colours_and_literal_filenames(self):
        from openpyxl import load_workbook
        self.write()
        result = self.check()
        result['filename'] = '=sample.tif'
        report = dict(scanner=str(self.root), checked_at='2026-10-04T12:00:00+00:00', results=[result])
        output = self.root / 'report.xlsx'
        validation.write_excel(report, output)
        wb = load_workbook(output)
        ws = wb.active
        self.assertEqual(ws['A8'].value, '=sample.tif')
        self.assertEqual(ws['A8'].data_type, 's')
        self.assertEqual(ws['C8'].value, 'PASS')
        self.assertEqual(ws['E8'].value, 64)
        self.assertEqual(ws.freeze_panes, 'D8')
        rules = list(ws.conditional_formatting)
        self.assertEqual(len(rules), 1)
        self.assertEqual(len(ws.conditional_formatting[rules[0]]), 3)
        wb.close()


if __name__ == '__main__':
    unittest.main()
