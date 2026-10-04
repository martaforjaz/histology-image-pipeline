"""Read-only validation of saved 2x/40x outputs, with an Excel status report.

Run without arguments to be prompted for the scanner directory. No conversion,
repair, deletion or renaming is performed. Each existing final file is checked
in a separate process so a broken decoder cannot hang the entire batch.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
import json
import logging
import math
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import time
import warnings
import xml.etree.ElementTree as ET

import numpy as np
import tifffile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '01_conversion_downsampling'))
from image_metadata import read_tiff_mpp, validate_mpp

RAW_EXTENSIONS = {'.ndpi', '.ndp', '.svs', '.scn', '.mrxs', '.qptiff', '.tif', '.tiff',
                  '.czi', '.vsi', '.dcm', '.isyntax', '.i2syntax'}
TARGETS = {'2x': 5.0, '40x': 0.25}
METHOD = ('PASS means complete TIFF decoding (all stored pages and pyramid levels), '
          'valid storage ranges, RGB layout, calibrated resolution and dimensions '
          'consistent with the raw source (1 pixel rounding tolerance). '
          'This is not a raw-to-output pixel equality or tissue-quality assessment. '
          'FAIL means a detected problem or missing/unfinished output. REVIEW means '
          'verification could not be completed. Missing outputs may not have been generated yet.')


def image_key(name):
    name = name.lower()
    if name.endswith('.part'):
        name = name[:-5]
    for suffix in ('.ome.tiff', '.ome.tif', *sorted(RAW_EXTENSIONS, key=len, reverse=True)):
        if name.endswith(suffix):
            return name[:-len(suffix)]
    return name


def inventory(scanner):
    """Include existing files and absent outputs for every identified raw scan."""
    sources = {}
    for path in sorted(scanner.iterdir()):
        if path.is_file() and path.suffix.lower() in RAW_EXTENSIONS:
            sources.setdefault(image_key(path.name), []).append(str(path))
    outputs = {}
    for folder in TARGETS:
        directory = scanner / folder
        if directory.exists():
            for path in sorted(directory.iterdir()):
                name = path.name.lower().removesuffix('.part')
                if path.is_file() and name.endswith(('.tif', '.tiff')):
                    outputs.setdefault((folder, image_key(path.name)), []).append(path)
    keys = sorted(set(sources) | {key for _, key in outputs})
    if not keys:
        raise ValueError('No supported raw scans or TIFF outputs found in the selected scanner folder.')
    jobs = []
    for key in keys:
        for folder, mpp in TARGETS.items():
            matches = outputs.get((folder, key), [])
            for path in matches or [None]:
                raw_name = Path(sources[key][0]).stem if key in sources else key
                missing_name = raw_name + ('.ome.tif' if folder == '40x' else '.tif')
                jobs.append(dict(image=key, folder=folder, filename=path.name if path else missing_name,
                                 path=str(path) if path else str(scanner / folder),
                                 source_paths=sources.get(key, []), expected_mpp=mpp,
                                 missing=path is None,
                                 duplicate=sum(not p.name.lower().endswith('.part') for p in matches) > 1))
    return jobs


def check_tiff_directories(path):
    """Reject truncated/cyclic directory graphs before asking a decoder to open."""
    size = path.stat().st_size
    with path.open('rb') as f:
        header = f.read(16)
        if len(header) < 8 or header[:2] not in (b'II', b'MM'):
            raise ValueError('Invalid or incomplete TIFF header')
        endian = '<' if header[:2] == b'II' else '>'
        version = struct.unpack(endian + 'H', header[2:4])[0]
        if version == 43:
            if len(header) < 16 or header[4:8] != struct.pack(endian + 'HH', 8, 0):
                raise ValueError('Invalid BigTIFF header')
            fmt, count_size, entry_size, value_size, offset = 'Q', 8, 20, 8, struct.unpack(endian+'Q', header[8:16])[0]
        elif version == 42:
            fmt, count_size, entry_size, value_size, offset = 'I', 2, 12, 4, struct.unpack(endian+'I', header[4:8])[0]
        else:
            raise ValueError('Not a TIFF or BigTIFF file')
        pending, seen = [offset], set()
        type_sizes = {1:1, 2:1, 3:2, 4:4, 5:8, 6:1, 7:1, 8:2, 9:4, 10:8, 11:4, 12:8, 13:4, 16:8, 17:8, 18:8}
        while pending:
            offset = pending.pop()
            if offset in seen:
                # SubIFDs can also form a next-IFD chain. Revisit is allowed but
                # never followed twice, preventing decoder loops in this scan.
                continue
            if offset < 8 or offset + count_size > size:
                raise ValueError(f'TIFF directory outside file at byte {offset}')
            seen.add(offset)
            if len(seen) > 256:
                raise ValueError('Unexpectedly many TIFF directories (>256)')
            f.seek(offset)
            count = struct.unpack(endian + ('Q' if count_size == 8 else 'H'), f.read(count_size))[0]
            if count > 4096 or offset + count_size + count*entry_size + value_size > size:
                raise ValueError('Truncated or invalid TIFF directory')
            entries = f.read(count*entry_size)
            next_offset = struct.unpack(endian+fmt, f.read(value_size))[0]
            if next_offset:
                if next_offset <= offset:
                    raise ValueError('Backward/cyclic TIFF directory pointer')
                pending.append(next_offset)
            for i in range(count):
                entry = entries[i*entry_size:(i+1)*entry_size]
                tag, typ = struct.unpack(endian+'HH', entry[:4])
                n = struct.unpack(endian+fmt, entry[4:4+value_size])[0]
                value = entry[-value_size:]
                length = n*type_sizes.get(typ, 1)
                pointer = struct.unpack(endian+fmt, value)[0]
                if length > value_size and (pointer < 8 or pointer+length > size):
                    raise ValueError(f'TIFF tag {tag} points beyond the end of the file')
                if tag == 330 and n:
                    if n > 64 or typ not in (4, 13, 16, 18):
                        raise ValueError('Invalid pyramid directory list')
                    if length > value_size:
                        f.seek(pointer)
                        value = f.read(length)
                    child_fmt = 'Q' if type_sizes[typ] == 8 else 'I'
                    children = struct.unpack(endian + str(n) + child_fmt, value[:length])
                    if any(child <= offset for child in children):
                        raise ValueError('Missing or cyclic pyramid directory pointer')
                    pending.extend(children)


def source_geometry(paths):
    if len(paths) != 1:
        raise ValueError('No unique matching raw source found')
    # OpenSlide supports the Hamamatsu and MRXS/P1000 sources requested here.
    # Other source formats are explicitly REVIEW if this backend cannot open them.
    from openslide import OpenSlide
    with OpenSlide(paths[0]) as slide:
        mx, my = validate_mpp(float(slide.properties['openslide.mpp-x']),
                              float(slide.properties['openslide.mpp-y']))
        return slide.dimensions, (mx, my)


class TiffMessages(logging.Handler):
    def __init__(self):
        super().__init__(logging.WARNING)
        self.messages = []

    def emit(self, record):
        self.messages.append(record.getMessage())


def decode_page(page, file_size, path):
    if page.dtype != np.dtype('uint8') or page.samplesperpixel != 3 or page.planarconfig != 1:
        raise ValueError('Expected an 8-bit RGB saved image')
    for offset, count in zip(page.dataoffsets, page.databytecounts):
        if offset < 8 or count <= 0 or offset+count > file_size:
            raise ValueError('Missing tile/strip or pixel data extends beyond file size')
    low, high = 255, 0
    last_progress = time.monotonic()
    # Uncompressed full-width strips can be hundreds of MB. Read their bytes
    # in small pieces; uint8 RGB with no predictor needs no codec transform.
    if page.compression == 1 and page.predictor == 1 and page.photometric == 2 and not page.is_tiled:
        expected_segments = math.ceil(page.imagelength / page.rowsperstrip)
        if len(page.dataoffsets) != expected_segments:
            raise ValueError('Incorrect strip count')
        with path.open('rb') as f:
            for i, (offset, count) in enumerate(zip(page.dataoffsets, page.databytecounts)):
                rows = min(page.rowsperstrip, page.imagelength-i*page.rowsperstrip)
                expected = rows*page.imagewidth*3
                if count != expected:
                    raise ValueError('Uncompressed strip byte count does not match dimensions')
                f.seek(offset)
                remaining = count
                while remaining:
                    block = f.read(min(16 << 20, remaining))
                    if not block:
                        raise ValueError('Unexpected end of pixel data')
                    values = np.frombuffer(block, dtype=np.uint8)
                    low, high = min(low, int(values.min())), max(high, int(values.max()))
                    remaining -= len(block)
                    if time.monotonic()-last_progress > 30:
                        print(f'  Reading {path.name}: strip {i+1}/{expected_segments}, {remaining:,} bytes left in strip', flush=True)
                        last_progress = time.monotonic()
    else:
        count = 0
        for segment, position, shape in page.segments(maxworkers=1, buffersize=4 << 20):
            if segment is None:
                raise ValueError('Undecodable or missing tile/strip')
            y, x = position[2:4]
            values = segment[0, :min(shape[1], page.imagelength-y), :min(shape[2], page.imagewidth-x)]
            if not values.size:
                raise ValueError('Empty decoded region')
            low, high = min(low, int(values.min())), max(high, int(values.max()))
            count += 1
            if time.monotonic()-last_progress > 30:
                print(f'  Decoding {path.name}: {page.imagewidth} x {page.imagelength}, tile/strip {count}/{len(page.dataoffsets)}', flush=True)
                last_progress = time.monotonic()
        if count != len(page.dataoffsets):
            raise ValueError('Not all tiles/strips were decoded')
    return low, high


def validate_file(job):
    started = time.monotonic()
    result = dict(job, status='FAIL', reason='', width=None, height=None,
                  mpp_x=None, mpp_y=None, pages_checked=0, full_decode=False,
                  source_check='NOT CHECKED', seconds=0)
    if job['missing']:
        result['reason'] = 'Missing output (may not have been generated yet)'
        return result
    path = Path(job['path'])
    if path.name.lower().endswith('.part'):
        result['reason'] = 'Unfinished .part file; final output has not been committed'
        return result
    handler = TiffMessages()
    logger = logging.getLogger('tifffile')
    logger.addHandler(handler)
    issues, reviews = [], []
    try:
        before = path.stat()
        check_tiff_directories(path)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            # Streamed TIFFs can interleave pyramid tiles. A larger read buffer
            # avoids a separate network request for every small tile even when
            # offsets are not contiguous; decoding remains tile by tile.
            with path.open('rb', buffering=16 << 20) as handle, tifffile.TiffFile(handle) as tif:
                main = tif.pages[0]
                result.update(width=main.imagewidth, height=main.imagelength)
                mx, my = read_tiff_mpp(path)
                result.update(mpp_x=mx, mpp_y=my)
                target = job['expected_mpp']
                try:
                    (sw, sh), (sx, sy) = source_geometry(job['source_paths'])
                    target = max(target, sx, sy)  # Match converter's no-upsampling policy.
                    expected = (math.ceil(sw/(target/sx)), math.ceil(sh/(target/sy)))
                    if abs(main.imagewidth-expected[0]) > 1 or abs(main.imagelength-expected[1]) > 1:
                        issues.append(f'Wrong dimensions: expected approximately {expected[0]} x {expected[1]} from raw source')
                        result['source_check'] = 'FAIL'
                    else:
                        result['source_check'] = 'PASS'
                except Exception as exc:
                    reviews.append(f'Raw-source geometry unavailable: {type(exc).__name__}: {exc}')
                    result['source_check'] = 'UNVERIFIED'
                if not all(math.isclose(value, target, rel_tol=1e-4, abs_tol=1e-6) for value in (mx, my)):
                    issues.append(f'Wrong pixel spacing: expected {target:g} um/pixel')
                if tif.ome_metadata:
                    pixels = ET.fromstring(tif.ome_metadata).find('.//{*}Pixels')
                    if pixels is None or (int(pixels.get('SizeX', '0')), int(pixels.get('SizeY', '0'))) != (main.imagewidth, main.imagelength):
                        issues.append('OME dimensions disagree with TIFF dimensions')
                if job['folder'] == '40x' and (not tif.is_ome or not main.pages):
                    issues.append('40x output is missing OME metadata or pyramid levels')
                seen = set()
                def decode(page, is_main=False):
                    if page.offset in seen:
                        raise ValueError('Duplicate/cyclic TIFF page')
                    seen.add(page.offset)
                    low, high = decode_page(page, before.st_size, path)
                    result['pages_checked'] += 1
                    if is_main and low == high:
                        reviews.append(f'All image values are {low}; image appears blank')
                    if page.pages:
                        previous = (page.imagewidth, page.imagelength)
                        for child in page.pages:
                            dims = (child.imagewidth, child.imagelength)
                            if not all(0 < a < b for a, b in zip(dims, previous)):
                                raise ValueError('Invalid pyramid dimensions/order')
                            decode(child)
                            previous = dims
                for i, page in enumerate(tif.pages):
                    decode(page, i == 0)
                result['full_decode'] = True
            reviews.extend(str(w.message) for w in caught)
        reviews.extend(handler.messages)
        after = path.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            issues.append('File changed during validation; rerun after saving finishes')
        if job['duplicate']:
            issues.append('Multiple final outputs match this image and resolution')
        result['status'] = 'FAIL' if issues else ('REVIEW' if reviews else 'PASS')
        result['reason'] = '; '.join(issues + reviews) or 'All integrity, resolution and source-dimension checks passed'
    except Exception as exc:
        if isinstance(exc, (PermissionError, FileNotFoundError, ImportError)):
            result['status'] = 'REVIEW'
        result['reason'] = f'{type(exc).__name__}: {exc}'
    finally:
        logger.removeHandler(handler)
        result['seconds'] = round(time.monotonic()-started, 2)
    return result


def run_isolated(job, timeout):
    if job['missing'] or job['filename'].lower().endswith('.part'):
        return validate_file(job)
    with tempfile.TemporaryDirectory(prefix='histology-validation-') as tmp:
        request, response = Path(tmp)/'request.json', Path(tmp)/'response.json'
        request.write_text(json.dumps(job), encoding='utf-8')
        try:
            process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()),
                                        '--worker', str(request), str(response)],
                                       stderr=subprocess.PIPE, text=True)
            try:
                _, stderr = process.communicate(timeout=timeout)
            except subprocess.TimeoutExpired:
                # Windows venv python.exe can launch a child interpreter. Kill
                # this worker's tree, not just the redirector, on timeout.
                if sys.platform == 'win32':
                    subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'],
                                   capture_output=True, timeout=15)
                else:
                    process.kill()
                process.communicate(timeout=15)
                raise
            if process.returncode != 0 or not response.exists():
                raise RuntimeError((stderr or f'Worker exited {process.returncode}')[-1000:])
            return json.loads(response.read_text(encoding='utf-8'))
        except (subprocess.TimeoutExpired, RuntimeError) as exc:
            reason = (f'Validation exceeded {timeout:g} seconds; retry with a longer --timeout or a local copy'
                      if isinstance(exc, subprocess.TimeoutExpired) else f'Validation process failed: {str(exc).splitlines()[-1]}')
            return dict(job, status='REVIEW', reason=reason,
                        full_decode=False, source_check='UNVERIFIED', seconds=timeout if isinstance(exc, subprocess.TimeoutExpired) else None)


def write_excel(report, path):
    """Portable Excel export for running this utility outside the Codex app."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import FormulaRule
    wb = Workbook()
    ws = wb.active
    ws.title = 'Validation'
    ws.sheet_view.showGridLines = False
    ws.append(['Downsampling validation'])
    ws.append(['Scanner folder', report['scanner']])
    ws.append(['Checked at (UTC)', report['checked_at']])
    ws.append(['Scope', METHOD])
    ws.append(['Legend', 'PASS = green; FAIL = red; REVIEW = amber (not fully verified)'])
    ws.append([])
    headers = ['File name', 'Folder', 'Status', 'Reason', 'Width (px)', 'Height (px)',
               'MPP X', 'MPP Y', 'Pages checked', 'Fully decoded', 'Source dimensions', 'Seconds', 'File path']
    ws.append(headers)
    keys = ['filename', 'folder', 'status', 'reason', 'width', 'height', 'mpp_x', 'mpp_y',
            'pages_checked', 'full_decode', 'source_check', 'seconds', 'path']
    for result in report['results']:
        ws.append([result.get(key) for key in keys])
    # Force strings to literal text, including filenames beginning with '='.
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='top')
    ws['A1'].font = Font(name='Arial', size=16, bold=True)
    for cell in ws[7]:
        cell.fill = PatternFill('solid', fgColor='243746')
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    for status, color in [('PASS', 'C6EFCE'), ('FAIL', 'FFC7CE'), ('REVIEW', 'FFEB9C')]:
        ws.conditional_formatting.add(f'A8:C{max(8,ws.max_row)}', FormulaRule(
            formula=[f'$C8="{status}"'], fill=PatternFill('solid', fgColor=color)))
    widths = {'A':38, 'B':15, 'C':14, 'D':90, 'E':15, 'F':15, 'G':14, 'H':14,
              'I':16, 'J':18, 'K':23, 'L':14, 'M':65}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in range(8, ws.max_row+1):
        ws.cell(row, 4).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row].height = 42
    ws.merge_cells('B4:M4')
    ws['B4'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[4].height = 48
    ws.freeze_panes = 'D8'
    ws.auto_filter.ref = f'A7:M{ws.max_row}'
    path = Path(path)
    temporary = path.with_suffix('.tmp.xlsx')
    wb.save(temporary)
    temporary.replace(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--scanner', type=Path, help='Scanner folder containing raw files and 2x/40x')
    parser.add_argument('--output', type=Path, help='Output .xlsx path (default: timestamped report in scanner folder)')
    parser.add_argument('--workers', type=int, default=2, help='Concurrent read-only file checks (default: 2)')
    parser.add_argument('--timeout', type=float, default=1800, help='Maximum seconds per existing file (default: 1800)')
    parser.add_argument('--json-only', action='store_true', help='Write validation results without Excel export')
    parser.add_argument('--report-from', type=Path, help='Create Excel from a completed JSON report without decoding again')
    parser.add_argument('--worker', nargs=2, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.output and args.output.suffix.lower() != '.xlsx':
        parser.error('--output must end in .xlsx')
    if args.worker:
        request, response = map(Path, args.worker)
        response.write_text(json.dumps(validate_file(json.loads(request.read_text(encoding='utf-8')))), encoding='utf-8')
        return
    if args.report_from:
        report = json.loads(args.report_from.read_text(encoding='utf-8'))
        if not report.get('complete'):
            parser.error('The JSON validation run is incomplete; wait for all file checks to finish.')
        output = args.output or args.report_from.with_suffix('.xlsx')
        if output.exists():
            parser.error('Report already exists; choose a different --output path.')
        output.parent.mkdir(parents=True, exist_ok=True)
        write_excel(report, output)
        print(f'Excel report: {output}')
        return
    if args.scanner is None:
        args.scanner = Path(input('Scanner folder containing 2x and 40x: ').strip().strip('"'))
    if not args.scanner.is_dir():
        parser.error(f'Scanner folder does not exist or is inaccessible: {args.scanner}')
    if args.workers < 1 or args.timeout <= 0:
        parser.error('Workers and timeout must be positive.')
    if not args.json_only:
        try:
            import openpyxl  # Fail before long decoding if Excel dependency is missing.
        except ImportError:
            parser.error('Install the report dependency first: python -m pip install openpyxl')
    now = datetime.now(timezone.utc)
    output = args.output or args.scanner / ('downsampling_validation_' + now.strftime('%Y%m%d_%H%M%S') + '.xlsx')
    output.parent.mkdir(parents=True, exist_ok=True)
    checkpoint = output.with_suffix('.json')
    if checkpoint.exists() or (not args.json_only and output.exists()):
        parser.error('Report already exists; choose a different --output path.')
    jobs = inventory(args.scanner)
    jobs.sort(key=lambda j: (not (j['missing'] or j['filename'].lower().endswith('.part')),
                             j['folder'] != '2x', j['image']))
    report = dict(scanner=str(args.scanner), checked_at=now.isoformat(), method=METHOD, complete=False,
                  expected_outputs=len(jobs), results=[])
    print(f'{len(jobs)} expected/existing outputs. Full decoding may take a long time over a network.', flush=True)
    def save_checkpoint():
        temporary = checkpoint.with_suffix('.json.tmp')
        temporary.write_text(json.dumps(report, indent=2), encoding='utf-8')
        temporary.replace(checkpoint)
    save_checkpoint()
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(run_isolated, job, args.timeout): job for job in jobs}
        for future in as_completed(futures):
            result = future.result()
            report['results'].append(result)
            save_checkpoint()
            print(f"[{len(report['results'])}/{len(jobs)}] {result['status']} {result['folder']} / {result['filename']}: {result['reason']}", flush=True)
    report['results'].sort(key=lambda r: (r['image'], r['folder'], r['filename']))
    report['complete'] = True
    save_checkpoint()
    if not args.json_only:
        write_excel(report, output)
        print(f'Excel report: {output}', flush=True)
    print('Counts: ' + json.dumps({s: sum(r['status']==s for r in report['results']) for s in ('PASS','FAIL','REVIEW')}), flush=True)
    print(f'Validation data: {checkpoint}', flush=True)


if __name__ == '__main__':
    main()
